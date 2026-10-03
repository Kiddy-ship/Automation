from datetime import date, datetime
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from automation_app.report_parser import EXPECTED_HEADERS, ReportParseError, parse_report


class OpenXmlReportTests(unittest.TestCase):
    def _write_compact_report(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Báo cáo"
        headers = [
            "",
            "",
            EXPECTED_HEADERS[7],
            EXPECTED_HEADERS[8],
            EXPECTED_HEADERS[12],
            EXPECTED_HEADERS[13],
            "Giá",
            "",
            EXPECTED_HEADERS[19],
            EXPECTED_HEADERS[15],
            "",
            "",
            EXPECTED_HEADERS[9],
            EXPECTED_HEADERS[10],
            EXPECTED_HEADERS[23],
            EXPECTED_HEADERS[24],
            EXPECTED_HEADERS[25],
        ]
        for column, value in enumerate(headers, 1):
            sheet.cell(row=1, column=column, value=value)
        row = [
            "",
            1,
            "DH-OPENXML-1",
            datetime(2026, 9, 11),
            "SKU-1",
            "Sản phẩm thử",
            100,
            "",
            0,
            2,
            "",
            "",
            "KH-1",
            "Khách thử",
            200,
            200,
            "Ghi chú",
        ]
        for column, value in enumerate(row, 1):
            sheet.cell(row=2, column=column, value=value)
        workbook.save(path)

    def test_xlsx_report_is_parsed(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            self._write_compact_report(path)
            report = parse_report(path, expected_date=date(2026, 9, 11))
            self.assertEqual(report.report_start_date, date(2026, 9, 11))
            self.assertEqual(len(report.orders), 1)
            self.assertEqual(report.orders[0].related_order_code, "DH-OPENXML-1")
            self.assertEqual(report.orders[0].items[0].quantity, 2)

    def test_xlsm_report_is_parsed(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsm"
            self._write_compact_report(path)
            report = parse_report(path, expected_date=date(2026, 9, 11))
            self.assertEqual(len(report.orders), 1)

    def test_unsupported_extension_has_actionable_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.csv"
            path.write_text("not excel", encoding="utf-8")
            with self.assertRaises(ReportParseError) as raised:
                parse_report(path)
            self.assertIn(".xlsx", str(raised.exception))


if __name__ == "__main__":
    unittest.main()
