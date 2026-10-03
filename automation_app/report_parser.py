"""Parse Sopoka's Excel sales-summary reports safely.

The exporter writes order-level values only on the first row of each order
(columns A:L) while product-level values are present on every row (M:U).
Columns V:Z can be vertically merged.  This module therefore treats every
non-empty value in column H as the start of a new order and forward-fills the
order context for its following product rows.

Both legacy BIFF ``.xls`` and OpenXML ``.xlsx``/``.xlsm`` workbooks are
supported.  Only read operations are performed; the source workbook is never
modified.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import math
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import Any, Iterable, Iterator, Mapping, Sequence, TypeAlias

import xlrd
from xlrd.sheet import Cell, Sheet


MAX_ORDERS_PER_BATCH = 19
EXPECTED_SHEET_NAME = "Báo cáo"
OLE2_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
OPENXML_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
SUPPORTED_REPORT_EXTENSIONS = frozenset({".xls", ".xlsx", ".xlsm", ".xltx", ".xltm"})
SUPPORTED_REPORT_FORMATS = ".xls, .xlsx, .xlsm, .xltx hoặc .xltm"
MONEY_TOLERANCE = Decimal("1")

EXPECTED_HEADERS: tuple[str, ...] = (
    "STT",
    "Ngày",
    "Mã nhân viên",
    "Tên nhân viên",
    "Số hóa đơn",
    "Ngày hóa đơn",
    "Nhân viên bán hàng",
    "Mã đặt đơn hàng liên quan",
    "Ngày đặt đơn hàng liên quan",
    "Mã khách",
    "Tên khách",
    "Địa chỉ",
    "Mã sản phẩm",
    "Tên sản phẩm",
    "Ngành hàng",
    "Số lượng bán",
    "Số lượng KM",
    "ĐVT",
    "Thành tiền (SL bán x giá)",
    "Chiết khấu (%) (từng SP)",
    "Tiền chiết khấu (từng SP)",
    "Chiết khấu (%) (tổng bill)",
    "Tiền chiết khấu (tổng bill)",
    "Thanh toán",
    "Doanh thu (-VAT)",
    "Diễn giải",
)

DateInput: TypeAlias = date | datetime | str


class _OpenXmlCell:
    """Small xlrd-compatible cell view backed by an openpyxl cell.

    The parser deliberately uses only a tiny subset of xlrd's cell API.  A
    compatibility view keeps the two workbook readers on the same validation
    path, so dates, identifiers, totals and layout checks behave identically
    for ``.xls`` and OpenXML workbooks.
    """

    __slots__ = ("ctype", "value", "xf_index", "number_format")

    def __init__(self, source: Any, *, epoch: Any) -> None:
        raw = source.value
        data_type = getattr(source, "data_type", None)
        self.xf_index = 0
        self.number_format = getattr(source, "number_format", "") or ""
        if raw is None:
            self.ctype = xlrd.XL_CELL_EMPTY
            self.value = ""
        elif data_type == "e":
            self.ctype = xlrd.XL_CELL_ERROR
            self.value = raw
        elif data_type == "d" and isinstance(raw, (int, float)):
            self.ctype = xlrd.XL_CELL_DATE
            self.value = raw
        elif isinstance(raw, (datetime, date)):
            # xlrd represents dates as serial numbers plus workbook.datemode.
            # openpyxl's conversion uses the workbook epoch and preserves the
            # same 1900/1904 date systems.
            try:
                from openpyxl.utils.datetime import to_excel

                self.value = to_excel(raw, epoch=epoch)
            except (TypeError, ValueError, OverflowError):
                self.ctype = xlrd.XL_CELL_TEXT
                self.value = raw.isoformat()
            else:
                self.ctype = xlrd.XL_CELL_DATE
        elif isinstance(raw, bool):
            self.ctype = xlrd.XL_CELL_BOOLEAN
            self.value = raw
        elif isinstance(raw, (int, float)):
            self.ctype = xlrd.XL_CELL_NUMBER
            self.value = raw
        else:
            self.ctype = xlrd.XL_CELL_TEXT
            self.value = raw


class _OpenXmlSheet:
    """xlrd-compatible worksheet view for a read-only openpyxl sheet."""

    __slots__ = ("_sheet", "_epoch", "name", "nrows", "ncols")

    def __init__(self, sheet: Any, *, epoch: Any) -> None:
        self._sheet = sheet
        self._epoch = epoch
        self.name = str(sheet.title)
        self.nrows = int(sheet.max_row or 0)
        self.ncols = int(sheet.max_column or 0)

    def cell(self, row: int, col: int) -> _OpenXmlCell:
        return _OpenXmlCell(
            self._sheet.cell(row=row + 1, column=col + 1), epoch=self._epoch
        )


class _OpenXmlWorkbook:
    """xlrd-compatible workbook view for ``.xlsx`` and ``.xlsm`` files."""

    __slots__ = ("_workbook", "_sheets", "datemode", "epoch")

    def __init__(self, workbook: Any) -> None:
        self._workbook = workbook
        self.epoch = getattr(workbook, "epoch", None)
        epoch_year = getattr(self.epoch, "year", 1899)
        self.datemode = 1 if epoch_year == 1904 else 0
        self._sheets = [
            _OpenXmlSheet(sheet, epoch=self.epoch)
            for sheet in workbook.worksheets
        ]

    def sheet_names(self) -> list[str]:
        return [sheet.name for sheet in self._sheets]

    def sheets(self) -> list[_OpenXmlSheet]:
        return list(self._sheets)

    def sheet_by_name(self, name: str) -> _OpenXmlSheet:
        for sheet in self._sheets:
            if sheet.name == name:
                return sheet
        raise KeyError(name)

    def release_resources(self) -> None:
        close = getattr(self._workbook, "close", None)
        if callable(close):
            close()


class _XmlCellSource:
    """Value/type carrier used by the dependency-free OpenXML fallback."""

    __slots__ = ("value", "data_type", "number_format")

    def __init__(self, value: Any = "", data_type: str | None = None, number_format: str = "") -> None:
        self.value = value
        self.data_type = data_type
        self.number_format = number_format


class _RawOpenXmlSheet:
    """Worksheet parsed directly from XML when openpyxl rejects bad styles."""

    __slots__ = ("name", "_cells", "nrows", "ncols")

    def __init__(self, name: str, cells: dict[tuple[int, int], _OpenXmlCell]) -> None:
        self.name = name
        self._cells = cells
        self.nrows = max((row for row, _column in cells), default=-1) + 1
        self.ncols = max((column for _row, column in cells), default=-1) + 1

    def cell(self, row: int, col: int) -> _OpenXmlCell:
        return self._cells.get((row, col), _OpenXmlCell(_XmlCellSource(), epoch=None))


class _RawOpenXmlWorkbook:
    """Minimal OpenXML reader that ignores cosmetic/style XML defects.

    Some Sopoka exports contain invalid legacy colour attributes such as
    ``rgb="13"``.  openpyxl correctly rejects those styles even though the
    cell values are readable.  This reader consumes only workbook, shared
    string and worksheet value XML, which is all the report parser needs.
    """

    __slots__ = ("_sheets", "datemode")

    def __init__(self, path: Path) -> None:
        with zipfile.ZipFile(path, "r") as archive:
            workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
            workbook_pr = workbook_root.find(f"{_MAIN_NS}workbookPr")
            date1904 = workbook_pr is not None and str(workbook_pr.attrib.get("date1904", "0")).casefold() in {"1", "true"}
            self.datemode = 1 if date1904 else 0
            formats = _read_openxml_number_formats(archive)
            shared_strings = _read_openxml_shared_strings(archive)
            relationships = _read_openxml_relationships(archive)
            sheets: list[_RawOpenXmlSheet] = []
            sheets_node = workbook_root.find(f"{_MAIN_NS}sheets")
            if sheets_node is None:
                raise ValueError("OpenXML workbook has no sheets")
            for sheet_node in sheets_node.findall(f"{_MAIN_NS}sheet"):
                name = str(sheet_node.attrib.get("name") or "")
                relationship_id = sheet_node.attrib.get(f"{_REL_NS}id")
                target = relationships.get(str(relationship_id or ""))
                if not target:
                    continue
                sheet_path = _openxml_target_path(target)
                cells = _read_openxml_sheet(
                    archive.read(sheet_path),
                    shared_strings=shared_strings,
                    formats=formats,
                )
                sheets.append(_RawOpenXmlSheet(name, cells))
            if not sheets:
                raise ValueError("OpenXML workbook has no readable worksheets")
            self._sheets = sheets

    def sheet_names(self) -> list[str]:
        return [sheet.name for sheet in self._sheets]

    def sheets(self) -> list[_RawOpenXmlSheet]:
        return list(self._sheets)

    def sheet_by_name(self, name: str) -> _RawOpenXmlSheet:
        for sheet in self._sheets:
            if sheet.name == name:
                return sheet
        raise KeyError(name)

    def release_resources(self) -> None:
        return None


_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PACKAGE_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _openxml_target_path(target: str) -> str:
    target = str(target).replace("\\", "/")
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join("xl", target))


def _read_openxml_relationships(archive: zipfile.ZipFile) -> dict[str, str]:
    path = "xl/_rels/workbook.xml.rels"
    if path not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(path))
    return {
        str(node.attrib.get("Id")): str(node.attrib.get("Target"))
        for node in root.findall(f"{_PACKAGE_REL_NS}Relationship")
        if node.attrib.get("Id") and node.attrib.get("Target")
    }


def _read_openxml_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    path = "xl/sharedStrings.xml"
    if path not in archive.namelist():
        return []
    root = ET.fromstring(archive.read(path))
    return [
        "".join(text.text or "" for text in node.iter(f"{_MAIN_NS}t"))
        for node in root.findall(f"{_MAIN_NS}si")
    ]


def _read_openxml_number_formats(archive: zipfile.ZipFile) -> dict[int, tuple[str, bool]]:
    path = "xl/styles.xml"
    if path not in archive.namelist():
        return {}
    try:
        root = ET.fromstring(archive.read(path))
    except ET.ParseError:
        # Styles are cosmetic for this parser.  A broken stylesheet should
        # not prevent reading otherwise valid cell values.
        return {}
    custom: dict[int, str] = {}
    numfmts = root.find(f"{_MAIN_NS}numFmts")
    if numfmts is not None:
        for node in numfmts.findall(f"{_MAIN_NS}numFmt"):
            try:
                custom[int(node.attrib.get("numFmtId", ""))] = str(node.attrib.get("formatCode") or "")
            except ValueError:
                continue
    # The built-in date IDs cover the formats most often emitted by Excel.
    builtin_dates = {14: "m/d/yy", 15: "d-mmm-yy", 16: "d-mmm", 17: "mmm-yy", 18: "h:mm AM/PM", 19: "h:mm:ss AM/PM", 20: "h:mm", 21: "h:mm:ss", 22: "m/d/yy h:mm", 45: "mm:ss", 46: "[h]:mm:ss", 47: "mmss.0"}
    result: dict[int, tuple[str, bool]] = {}
    cell_xfs = root.find(f"{_MAIN_NS}cellXfs")
    if cell_xfs is None:
        return result
    for index, xf in enumerate(cell_xfs.findall(f"{_MAIN_NS}xf")):
        try:
            num_fmt_id = int(xf.attrib.get("numFmtId", "0"))
        except ValueError:
            num_fmt_id = 0
        format_string = custom.get(num_fmt_id, builtin_dates.get(num_fmt_id, ""))
        result[index] = (format_string, _is_openxml_date_format(format_string, num_fmt_id))
    return result


def _is_openxml_date_format(format_string: str, num_fmt_id: int = 0) -> bool:
    if num_fmt_id in {14, 15, 16, 17, 18, 19, 20, 21, 22, 45, 46, 47}:
        return True
    cleaned = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", str(format_string)).casefold()
    return bool(re.search(r"[dyhms]", cleaned)) and not bool(re.search(r"[^0#.,/ :;dyhms]", cleaned))


def _column_index_from_reference(reference: str) -> int:
    letters = re.match(r"([A-Za-z]+)", reference or "")
    if not letters:
        return 0
    value = 0
    for character in letters.group(1).upper():
        value = value * 26 + ord(character) - ord("A") + 1
    return value - 1


def _read_openxml_sheet(
    payload: bytes,
    *,
    shared_strings: Sequence[str],
    formats: Mapping[int, tuple[str, bool]],
) -> dict[tuple[int, int], _OpenXmlCell]:
    root = ET.fromstring(payload)
    cells: dict[tuple[int, int], _OpenXmlCell] = {}
    sheet_data = root.find(f"{_MAIN_NS}sheetData")
    if sheet_data is None:
        return cells
    for row_index, row_node in enumerate(sheet_data.findall(f"{_MAIN_NS}row")):
        try:
            row = int(row_node.attrib.get("r", str(row_index + 1))) - 1
        except ValueError:
            row = row_index
        fallback_column = 0
        for cell_node in row_node.findall(f"{_MAIN_NS}c"):
            reference = str(cell_node.attrib.get("r") or "")
            column = _column_index_from_reference(reference) if reference else fallback_column
            fallback_column = column + 1
            try:
                style_index = int(cell_node.attrib.get("s", "0"))
            except ValueError:
                style_index = 0
            format_string, is_date = formats.get(style_index, ("", False))
            value_type = str(cell_node.attrib.get("t") or "")
            value_node = cell_node.find(f"{_MAIN_NS}v")
            raw_text = "" if value_node is None or value_node.text is None else value_node.text
            if value_type == "inlineStr":
                value = "".join(text.text or "" for text in cell_node.iter(f"{_MAIN_NS}t"))
                data_type = None
            elif value_type == "s":
                try:
                    value = shared_strings[int(raw_text)]
                except (ValueError, IndexError):
                    value = ""
                data_type = None
            elif value_type in {"str", "e"}:
                value = raw_text
                data_type = value_type
            elif value_type == "b":
                value = raw_text in {"1", "true", "TRUE"}
                data_type = "b"
            elif raw_text == "":
                value = ""
                data_type = None
            else:
                try:
                    value = float(raw_text)
                    if value.is_integer() and "." not in raw_text and "e" not in raw_text.casefold():
                        value = int(value)
                except ValueError:
                    value = raw_text
                data_type = "d" if is_date and isinstance(value, (int, float)) else None
            cells[(row, column)] = _OpenXmlCell(
                _XmlCellSource(value, data_type=data_type, number_format=format_string),
                epoch=None,
            )
    return cells


def _open_report_workbook(path: Path) -> Any:
    """Open a validated report with the reader appropriate for its suffix."""

    if path.suffix.casefold() == ".xls":
        return xlrd.open_workbook(
            filename=str(path),
            formatting_info=True,
            on_demand=True,
            use_mmap=False,
        )

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - guarded by requirements.txt
        raise ReportParseError(
            "Chưa cài bộ đọc Excel OpenXML cho file báo cáo.",
            hint="Cài openpyxl rồi thử lại.",
        ) from exc
    try:
        workbook = load_workbook(
            filename=str(path),
            read_only=True,
            data_only=True,
            keep_vba=path.suffix.casefold() in {".xlsm", ".xltm"},
        )
    except Exception:
        # A malformed cosmetic style must not make a value-only Sopoka export
        # unusable.  Fall back to the narrow XML reader, which intentionally
        # ignores styles while retaining values, dates and shared strings.
        return _RawOpenXmlWorkbook(path)
    return _OpenXmlWorkbook(workbook)


class ReportParseError(ValueError):
    """A user-facing error raised when a report is invalid or unsupported."""

    def __init__(
        self,
        message: str,
        *,
        cell: str | None = None,
        hint: str | None = None,
    ) -> None:
        self.message = message
        self.cell = cell
        self.hint = hint
        parts = [message]
        if cell:
            parts.append(f"Ô liên quan: {cell}.")
        if hint:
            parts.append(f"Gợi ý: {hint}")
        super().__init__(" ".join(parts))


@dataclass(slots=True)
class OrderItem:
    """One product line belonging to an order."""

    source_row: int
    product_code: str
    product_name: str
    category: str
    quantity: Decimal
    promotion_quantity: Decimal
    unit: str
    gross_amount: Decimal
    unit_price: Decimal
    discount_percent: Decimal
    discount_amount: Decimal
    net_amount: Decimal

    def to_dict(self) -> dict[str, Any]:
        """Return a JSON-serializable representation."""

        return _make_json_serializable(asdict(self))


@dataclass(slots=True)
class OrderRecord:
    """One related order from column H and all of its product lines."""

    sequence: int
    source_start_row: int
    source_end_row: int
    sale_date: date
    employee_code: str
    employee_name: str
    invoice_number: str
    invoice_date: date
    salesperson: str
    related_order_code: str
    related_order_date: date
    customer_code: str
    customer_name: str
    customer_address: str
    bill_discount_percent: Decimal | None
    bill_discount_amount: Decimal | None
    bill_discount_percent_error: str | None
    bill_discount_amount_error: str | None
    payment_amount: Decimal
    revenue_ex_vat: Decimal
    note: str
    items: list[OrderItem] = field(default_factory=list)

    @property
    def gross_total(self) -> Decimal:
        """Total gross value of all product lines."""

        return sum((item.gross_amount for item in self.items), Decimal("0"))

    @property
    def line_discount_total(self) -> Decimal:
        """Total line-level discount value."""

        return sum((item.discount_amount for item in self.items), Decimal("0"))

    @property
    def calculated_net_total(self) -> Decimal:
        """Net value calculated from product lines."""

        return sum((item.net_amount for item in self.items), Decimal("0"))

    def to_dict(self) -> dict[str, Any]:
        """Return a JSON-serializable order dictionary."""

        result = _make_json_serializable(asdict(self))
        result["gross_total"] = _json_decimal(self.gross_total)
        result["line_discount_total"] = _json_decimal(self.line_discount_total)
        result["calculated_net_total"] = _json_decimal(self.calculated_net_total)
        return result


@dataclass(slots=True)
class OrderBatch:
    """A KiotViet import batch containing at most 19 orders."""

    batch_number: int
    orders: list[OrderRecord]

    @property
    def order_count(self) -> int:
        """Number of orders/sheets in the batch."""

        return len(self.orders)

    @property
    def line_count(self) -> int:
        """Number of product lines across the batch."""

        return sum(len(order.items) for order in self.orders)

    def to_dict(self) -> dict[str, Any]:
        """Return a JSON-serializable batch dictionary."""

        return {
            "batch_number": self.batch_number,
            "order_count": self.order_count,
            "line_count": self.line_count,
            "orders": [order.to_dict() for order in self.orders],
        }


@dataclass(slots=True)
class ReportQA:
    """Reconciled totals and validation metadata for the parsed report."""

    order_count: int
    line_count: int
    gross_total: Decimal
    line_discount_total: Decimal
    calculated_net_total: Decimal
    payment_total: Decimal
    revenue_ex_vat_total: Decimal
    batch_count: int
    checks: dict[str, bool]

    def to_dict(self) -> dict[str, Any]:
        """Return JSON-safe QA data."""

        return _make_json_serializable(asdict(self))


@dataclass(slots=True)
class ParsedReport:
    """Structured, validated content of a Sopoka report."""

    source_file: str
    sheet_name: str
    header_row: int
    first_data_row: int
    last_data_row: int
    report_start_date: date
    report_end_date: date
    orders: list[OrderRecord]
    qa: ReportQA

    def batches(self, max_orders: int = MAX_ORDERS_PER_BATCH) -> list[OrderBatch]:
        """Split orders into deterministic batches, preserving source order."""

        return batch_orders(self.orders, max_orders=max_orders)

    def to_dict(
        self,
        *,
        include_batches: bool = True,
        max_orders_per_batch: int = MAX_ORDERS_PER_BATCH,
    ) -> dict[str, Any]:
        """Return a JSON-serializable report dictionary."""

        result: dict[str, Any] = {
            "source_file": self.source_file,
            "sheet_name": self.sheet_name,
            "header_row": self.header_row,
            "first_data_row": self.first_data_row,
            "last_data_row": self.last_data_row,
            "report_start_date": self.report_start_date.isoformat(),
            "report_end_date": self.report_end_date.isoformat(),
            "orders": [order.to_dict() for order in self.orders],
            "qa": self.qa.to_dict(),
        }
        if include_batches:
            result["batches"] = [
                batch.to_dict()
                for batch in self.batches(max_orders=max_orders_per_batch)
            ]
        return result


def parse_report(
    file_path: str | Path,
    *,
    expected_date: DateInput | None = None,
    sheet_name: str = EXPECTED_SHEET_NAME,
    max_orders_per_batch: int = MAX_ORDERS_PER_BATCH,
    require_single_day: bool = True,
) -> ParsedReport:
    """Parse and validate a Sopoka sales-summary report.

    Args:
        file_path: Source Excel path.  Legacy ``.xls`` is opened by xlrd;
            OpenXML files are opened read-only by openpyxl.
        expected_date: Optional business date.  Pass ``date.today()`` in the
            live workflow to reject a stale download.  Both ``dd/mm/YYYY`` and
            ISO ``YYYY-mm-dd`` strings are accepted.
        sheet_name: Preferred worksheet name.
        max_orders_per_batch: Batch size used by QA metadata; KiotViet allows
            no more than 19 sheets/orders per import.
        require_single_day: Reject reports whose title covers multiple dates.

    Returns:
        A :class:`ParsedReport` containing dataclasses and JSON-safe helpers.

    Raises:
        ReportParseError: If the workbook layout, data, date, or totals fail
            validation.
    """

    path = _validate_source_file(file_path)
    expected = _coerce_expected_date(expected_date) if expected_date else None
    if not 1 <= max_orders_per_batch <= MAX_ORDERS_PER_BATCH:
        raise ReportParseError(
            "Kích thước batch không hợp lệ.",
            hint=f"Chọn từ 1 đến {MAX_ORDERS_PER_BATCH} đơn mỗi batch.",
        )

    try:
        workbook = _open_report_workbook(path)
    except PermissionError as exc:
        raise ReportParseError(
            "Không thể đọc file báo cáo vì Windows đang chặn quyền truy cập.",
            hint="Đóng hộp thoại lưu file nếu đang mở và thử lại.",
        ) from exc
    except (xlrd.XLRDError, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise ReportParseError(
            "Không thể mở báo cáo Excel.",
            hint="Hãy xuất lại file từ Sopoka và giữ đúng đuôi .xls/.xlsx/.xlsm.",
        ) from exc
    except ReportParseError:
        raise
    except Exception as exc:
        # openpyxl exposes a few format-specific exceptions that are not
        # shared with xlrd (for example InvalidFileException).  Keep those
        # details out of the UI while returning the same actionable message.
        raise ReportParseError(
            "Không thể mở báo cáo Excel.",
            hint="Hãy xuất lại file từ Sopoka và giữ đúng đuôi .xls/.xlsx/.xlsm.",
        ) from exc

    try:
        sheet = _select_sheet(workbook, sheet_name)
        header_index, layout = _detect_report_layout(sheet)
        report_start, report_end = (
            _read_report_period(sheet, header_index)
            if layout == "full"
            else _read_compact_report_period(workbook, sheet, header_index)
        )
        _validate_report_period(
            report_start,
            report_end,
            expected_date=expected,
            require_single_day=require_single_day,
        )
        last_data_index = _find_last_data_row(
            sheet, header_index, product_column=12 if layout == "full" else 4
        )
        orders = (
            _parse_orders(
                workbook,
                sheet,
                header_index=header_index,
                last_data_index=last_data_index,
                report_start=report_start,
                report_end=report_end,
            )
            if layout == "full"
            else _parse_compact_orders(
                workbook,
                sheet,
                header_index=header_index,
                last_data_index=last_data_index,
                report_start=report_start,
                report_end=report_end,
            )
        )
        _ensure_unique_business_keys(orders)
        batches = batch_orders(orders, max_orders=max_orders_per_batch)
        qa = _build_qa(orders, batch_count=len(batches))
        return ParsedReport(
            source_file=str(path),
            sheet_name=sheet.name,
            header_row=header_index + 1,
            first_data_row=header_index + 2,
            last_data_row=last_data_index + 1,
            report_start_date=report_start,
            report_end_date=report_end,
            orders=orders,
            qa=qa,
        )
    finally:
        release = getattr(workbook, "release_resources", None)
        if callable(release):
            release()


def parse_today_report(
    file_path: str | Path,
    *,
    sheet_name: str = EXPECTED_SHEET_NAME,
    max_orders_per_batch: int = MAX_ORDERS_PER_BATCH,
) -> ParsedReport:
    """Parse a report and require its report period to be today's date."""

    return parse_report(
        file_path,
        expected_date=date.today(),
        sheet_name=sheet_name,
        max_orders_per_batch=max_orders_per_batch,
        require_single_day=True,
    )


def parse_report_to_dict(
    file_path: str | Path,
    *,
    expected_date: DateInput | None = None,
    max_orders_per_batch: int = MAX_ORDERS_PER_BATCH,
) -> dict[str, Any]:
    """Parse a report and return a directly JSON-serializable dictionary."""

    report = parse_report(
        file_path,
        expected_date=expected_date,
        max_orders_per_batch=max_orders_per_batch,
    )
    return report.to_dict(max_orders_per_batch=max_orders_per_batch)


def batch_orders(
    orders: Sequence[OrderRecord] | Iterable[OrderRecord],
    *,
    max_orders: int = MAX_ORDERS_PER_BATCH,
) -> list[OrderBatch]:
    """Split orders into batches no larger than KiotViet's 19-sheet limit."""

    if not 1 <= max_orders <= MAX_ORDERS_PER_BATCH:
        raise ReportParseError(
            "Số đơn tối đa trong một batch không hợp lệ.",
            hint=f"Giá trị hợp lệ là 1..{MAX_ORDERS_PER_BATCH}.",
        )
    materialized = list(orders)
    return [
        OrderBatch(
            batch_number=(start // max_orders) + 1,
            orders=materialized[start : start + max_orders],
        )
        for start in range(0, len(materialized), max_orders)
    ]


def _validate_source_file(file_path: str | Path) -> Path:
    path = Path(file_path).expanduser().resolve()
    if not path.exists():
        raise ReportParseError(
            f"Không tìm thấy file báo cáo: {path}",
            hint="Hãy tải lại báo cáo Excel từ Sopoka.",
        )
    if not path.is_file():
        raise ReportParseError("Đường dẫn báo cáo không phải là một file.")
    suffix = path.suffix.casefold()
    if suffix not in SUPPORTED_REPORT_EXTENSIONS:
        raise ReportParseError(
            f"File nguồn phải có định dạng Excel ({SUPPORTED_REPORT_FORMATS}).",
            hint="Dùng file report do Sopoka xuất ra và giữ nguyên đuôi file.",
        )
    try:
        with path.open("rb") as stream:
            signature = stream.read(max(len(OLE2_SIGNATURE), 4))
    except OSError as exc:
        raise ReportParseError("Không thể đọc file báo cáo.") from exc
    if suffix == ".xls" and signature[: len(OLE2_SIGNATURE)] != OLE2_SIGNATURE:
        raise ReportParseError(
            "File có đuôi .xls nhưng không phải workbook BIFF hợp lệ.",
            hint="Không dùng file HTML hoặc XLSX đã đổi đuôi thành .xls.",
        )
    if suffix != ".xls" and signature[:4] not in OPENXML_SIGNATURES:
        raise ReportParseError(
            f"File có đuôi {suffix} nhưng không phải workbook OpenXML hợp lệ.",
            hint="Không đổi đuôi thủ công; hãy xuất lại đúng định dạng Excel.",
        )
    if suffix != ".xls" and not zipfile.is_zipfile(path):
        raise ReportParseError(
            f"File có đuôi {suffix} nhưng không phải file Excel OpenXML hợp lệ.",
            hint="Hãy mở file bằng Excel và lưu lại dưới dạng .xlsx hoặc .xlsm.",
        )
    return path


def _select_sheet(workbook: xlrd.book.Book, preferred_name: str) -> Sheet:
    if preferred_name in workbook.sheet_names():
        return workbook.sheet_by_name(preferred_name)
    for sheet in workbook.sheets():
        try:
            _detect_report_layout(sheet)
        except ReportParseError:
            continue
        return sheet
    available = ", ".join(workbook.sheet_names()) or "(không có sheet)"
    raise ReportParseError(
        f"Không tìm thấy sheet báo cáo hợp lệ. Các sheet hiện có: {available}.",
        hint=f"Sheet mong đợi là '{preferred_name}'.",
    )


def _detect_report_layout(sheet: Sheet) -> tuple[int, str]:
    """Return the header row and one of Sopoka's two observed XLS layouts."""

    try:
        return _find_and_validate_header(sheet), "full"
    except ReportParseError as full_error:
        search_limit = min(sheet.nrows, 25)
        for row_index in range(search_limit):
            if sheet.ncols < 17:
                continue
            values = [_normalize_header(_cell_text(sheet.cell(row_index, column))) for column in range(17)]
            if (
                values[2] == _normalize_header("Mã đặt đơn hàng liên quan")
                and values[3] == _normalize_header("Ngày đặt đơn hàng liên quan")
                and values[4] == _normalize_header("Mã sản phẩm")
                and values[5] == _normalize_header("Tên sản phẩm")
                and values[6] in {_normalize_header("Giá"), _normalize_header("Đơn giá")}
                and values[8] == _normalize_header("Chiết khấu (%) (từng SP)")
                and values[9] == _normalize_header("Số lượng bán")
                and values[12] == _normalize_header("Mã khách")
                and values[13] == _normalize_header("Tên khách")
                and values[14] == _normalize_header("Thanh toán")
                and values[15] == _normalize_header("Doanh thu (-VAT)")
            ):
                return row_index, "compact"
        raise full_error


def _find_and_validate_header(sheet: Sheet) -> int:
    expected_h = _normalize_header(EXPECTED_HEADERS[7])
    expected_m = _normalize_header(EXPECTED_HEADERS[12])
    search_limit = min(sheet.nrows, 25)
    header_index: int | None = None
    for row_index in range(search_limit):
        if sheet.ncols < len(EXPECTED_HEADERS):
            break
        h_value = _normalize_header(_cell_text(sheet.cell(row_index, 7)))
        m_value = _normalize_header(_cell_text(sheet.cell(row_index, 12)))
        if h_value == expected_h and m_value == expected_m:
            header_index = row_index
            break
    if header_index is None:
        raise ReportParseError(
            "Không tìm thấy dòng tiêu đề của báo cáo bán hàng.",
            hint="Báo cáo phải có cột 'Mã đặt đơn hàng liên quan' tại cột H.",
        )

    mismatches: list[str] = []
    for column, expected in enumerate(EXPECTED_HEADERS):
        actual = _cell_text(sheet.cell(header_index, column))
        if _normalize_header(actual) != _normalize_header(expected):
            mismatches.append(
                f"{_column_letter(column)}: mong đợi '{expected}', nhận '{actual}'"
            )
    if mismatches:
        preview = "; ".join(mismatches[:5])
        if len(mismatches) > 5:
            preview += f"; và {len(mismatches) - 5} cột khác"
        raise ReportParseError(
            f"Cấu trúc cột của báo cáo đã thay đổi: {preview}.",
            cell=f"A{header_index + 1}:Z{header_index + 1}",
            hint="Xuất lại đúng báo cáo Tổng hợp bán hàng hoặc cập nhật mapping.",
        )
    return header_index


def _read_report_period(sheet: Sheet, header_index: int) -> tuple[date, date]:
    text_parts: list[str] = []
    for row_index in range(header_index):
        for column in range(min(sheet.ncols, len(EXPECTED_HEADERS))):
            value = _cell_text(sheet.cell(row_index, column))
            if value:
                text_parts.append(value)
    title_text = " ".join(text_parts)
    match = re.search(
        r"Từ\s*ngày\s*:\s*(\d{1,2}/\d{1,2}/\d{4})"
        r"\s*đến\s*ngày\s*:\s*(\d{1,2}/\d{1,2}/\d{4})",
        title_text,
        flags=re.IGNORECASE,
    )
    if not match:
        raise ReportParseError(
            "Không đọc được khoảng ngày trong tiêu đề báo cáo.",
            hint="Tiêu đề cần có dạng 'Từ ngày: dd/mm/YYYY đến ngày: dd/mm/YYYY'.",
        )
    return (
        _parse_date_text(match.group(1), field_name="Từ ngày"),
        _parse_date_text(match.group(2), field_name="Đến ngày"),
    )


def _read_compact_report_period(
    workbook: xlrd.book.Book, sheet: Sheet, header_index: int
) -> tuple[date, date]:
    dates: list[date] = []
    for row_index in range(header_index + 1, sheet.nrows):
        if not _cell_text(sheet.cell(row_index, 4)):
            continue
        if not _identifier(
            workbook,
            sheet.cell(row_index, 2),
            field_name="Mã đặt đơn hàng liên quan",
        ):
            continue
        dates.append(
            _date_from_cell(
                workbook,
                sheet.cell(row_index, 3),
                field_name="Ngày đặt đơn hàng liên quan",
                excel_row=row_index + 1,
            )
        )
    if not dates:
        raise ReportParseError(
            "Không đọc được ngày trong báo cáo rút gọn.",
            hint="Hãy chọn Hôm nay, bấm Tìm kiếm rồi Xuất Excel lại.",
        )
    return min(dates), max(dates)


def _validate_report_period(
    start: date,
    end: date,
    *,
    expected_date: date | None,
    require_single_day: bool,
) -> None:
    if start > end:
        raise ReportParseError("Ngày bắt đầu báo cáo lớn hơn ngày kết thúc.")
    if require_single_day and start != end:
        raise ReportParseError(
            f"Báo cáo đang chứa nhiều ngày ({_format_date(start)} - {_format_date(end)}).",
            hint="Trên Sopoka hãy chọn 'Hôm nay' rồi xuất lại Excel.",
        )
    if expected_date is not None and (start != expected_date or end != expected_date):
        raise ReportParseError(
            f"Báo cáo không đúng ngày yêu cầu. File đang là "
            f"{_format_date(start)} - {_format_date(end)}, cần {_format_date(expected_date)}.",
            hint="Chọn 'Hôm nay', bấm Áp dụng, Tìm kiếm và Xuất Excel lại.",
        )


def _find_last_data_row(
    sheet: Sheet, header_index: int, *, product_column: int = 12
) -> int:
    # Column M is mandatory for every product line.  Using it avoids the
    # formatted-but-empty rows commonly included at the end of Sopoka files.
    product_rows = [
        row_index
        for row_index in range(header_index + 1, sheet.nrows)
        if _cell_text(sheet.cell(row_index, product_column))
    ]
    if not product_rows:
        raise ReportParseError(
            "Báo cáo không có dòng sản phẩm nào.",
            hint="Kiểm tra bộ lọc ngày và bấm Tìm kiếm trước khi xuất Excel.",
        )
    return product_rows[-1]


def _parse_compact_orders(
    workbook: xlrd.book.Book,
    sheet: Sheet,
    *,
    header_index: int,
    last_data_index: int,
    report_start: date,
    report_end: date,
) -> list[OrderRecord]:
    """Parse Sopoka's grouped 17-column export without total-label rows."""

    orders: list[OrderRecord] = []
    by_code: dict[str, OrderRecord] = {}
    for row_index in range(header_index + 1, last_data_index + 1):
        excel_row = row_index + 1
        product_code = _identifier(
            workbook, sheet.cell(row_index, 4), field_name="Mã sản phẩm"
        )
        if not product_code:
            # Group footer rows such as "1 Total" are not product data.
            continue
        order_code = _identifier(
            workbook,
            sheet.cell(row_index, 2),
            field_name="Mã đặt đơn hàng liên quan",
        )
        if not order_code:
            raise ReportParseError(
                "Dòng sản phẩm thiếu mã đặt đơn hàng liên quan.",
                cell=f"C{excel_row}",
            )

        related_date = _date_from_cell(
            workbook,
            sheet.cell(row_index, 3),
            field_name="Ngày đặt đơn hàng liên quan",
            excel_row=excel_row,
        )
        if not report_start <= related_date <= report_end:
            raise ReportParseError(
                "Ngày đặt đơn hàng nằm ngoài khoảng báo cáo.", cell=f"D{excel_row}"
            )

        customer_code = _identifier(
            workbook, sheet.cell(row_index, 12), field_name="Mã khách"
        )
        customer_name = _clean_display_text(_cell_text(sheet.cell(row_index, 13)))
        if not customer_code or not customer_name:
            raise ReportParseError(
                "Dòng sản phẩm thiếu thông tin khách hàng.", cell=f"M{excel_row}:N{excel_row}"
            )

        order = by_code.get(order_code)
        if order is None:
            sequence = _required_decimal(
                sheet.cell(row_index, 1), field_name="STT", excel_row=excel_row,
                column_index=1,
            )
            if sequence != sequence.to_integral_value():
                raise ReportParseError("STT phải là số nguyên.", cell=f"B{excel_row}")
            order = OrderRecord(
                sequence=int(sequence),
                source_start_row=excel_row,
                source_end_row=excel_row,
                sale_date=related_date,
                employee_code="",
                employee_name="",
                invoice_number=f"compact:{order_code}",
                invoice_date=related_date,
                salesperson="",
                related_order_code=order_code,
                related_order_date=related_date,
                customer_code=customer_code,
                customer_name=customer_name,
                customer_address="",
                bill_discount_percent=None,
                bill_discount_amount=None,
                bill_discount_percent_error=None,
                bill_discount_amount_error=None,
                payment_amount=_required_decimal(
                    sheet.cell(row_index, 14),
                    field_name="Thanh toán",
                    excel_row=excel_row,
                    column_index=14,
                ),
                revenue_ex_vat=_required_decimal(
                    sheet.cell(row_index, 15),
                    field_name="Doanh thu (-VAT)",
                    excel_row=excel_row,
                    column_index=15,
                ),
                note=_clean_display_text(_cell_text(sheet.cell(row_index, 16))),
            )
            by_code[order_code] = order
            orders.append(order)
        elif (
            order.related_order_date != related_date
            or order.customer_code != customer_code
            or order.customer_name != customer_name
        ):
            raise ReportParseError(
                f"Thông tin đơn {order_code} thay đổi giữa các dòng sản phẩm.",
                cell=f"C{excel_row}:N{excel_row}",
            )

        product_name = _clean_display_text(_cell_text(sheet.cell(row_index, 5)))
        if not product_name:
            raise ReportParseError("Dòng sản phẩm thiếu tên sản phẩm.", cell=f"F{excel_row}")
        unit_price = _required_decimal(
            sheet.cell(row_index, 6), field_name="Giá", excel_row=excel_row,
            column_index=6,
        )
        discount_percent = _required_decimal(
            sheet.cell(row_index, 8),
            field_name="Chiết khấu (%) (từng SP)",
            excel_row=excel_row,
            column_index=8,
        )
        quantity = _required_decimal(
            sheet.cell(row_index, 9), field_name="Số lượng bán", excel_row=excel_row,
            column_index=9,
        )
        if unit_price < 0 or quantity <= 0:
            raise ReportParseError("Giá hoặc số lượng sản phẩm không hợp lệ.", cell=f"G{excel_row}:J{excel_row}")
        if not Decimal("0") <= discount_percent <= Decimal("100"):
            raise ReportParseError("Phần trăm chiết khấu không hợp lệ.", cell=f"I{excel_row}")
        gross_amount = unit_price * quantity
        discount_amount = gross_amount * discount_percent / Decimal("100")
        order.items.append(
            OrderItem(
                source_row=excel_row,
                product_code=product_code,
                product_name=product_name,
                category="",
                quantity=quantity,
                promotion_quantity=Decimal("0"),
                unit="",
                gross_amount=gross_amount,
                unit_price=unit_price,
                discount_percent=discount_percent,
                discount_amount=discount_amount,
                net_amount=gross_amount - discount_amount,
            )
        )
        order.source_end_row = excel_row

    if not orders:
        raise ReportParseError("Không tìm thấy đơn hàng hợp lệ trong báo cáo.")
    for order in orders:
        _validate_order_totals(order)
    return orders


def _parse_orders(
    workbook: xlrd.book.Book,
    sheet: Sheet,
    *,
    header_index: int,
    last_data_index: int,
    report_start: date,
    report_end: date,
) -> list[OrderRecord]:
    orders: list[OrderRecord] = []
    current: OrderRecord | None = None
    forward_context: list[str] | None = None

    for row_index in range(header_index + 1, last_data_index + 1):
        excel_row = row_index + 1
        product_code = _identifier(
            workbook, sheet.cell(row_index, 12), field_name="Mã sản phẩm"
        )
        related_order_raw = _identifier(
            workbook,
            sheet.cell(row_index, 7),
            field_name="Mã đặt đơn hàng liên quan",
        )

        if not product_code:
            if _row_has_any_value(sheet, row_index, 26):
                raise ReportParseError(
                    "Dòng dữ liệu không có mã sản phẩm.",
                    cell=f"M{excel_row}",
                )
            continue

        raw_context = [
            _identifier(workbook, sheet.cell(row_index, column), field_name=EXPECTED_HEADERS[column])
            for column in range(12)
        ]

        if related_order_raw:
            if current is not None:
                _validate_order_totals(current)
                orders.append(current)
            forward_context = raw_context
            current = _new_order_from_row(
                workbook,
                sheet,
                row_index=row_index,
                context=forward_context,
                report_start=report_start,
                report_end=report_end,
            )
        elif current is None or forward_context is None:
            raise ReportParseError(
                "Dòng sản phẩm xuất hiện trước mã đặt đơn hàng đầu tiên.",
                cell=f"H{excel_row}",
                hint="Mỗi nhóm phải bắt đầu bằng một mã ở cột H.",
            )
        else:
            forward_context = _forward_fill_and_validate_context(
                raw_context,
                forward_context,
                row_number=excel_row,
            )

        assert current is not None
        current.items.append(_parse_item(workbook, sheet, row_index))
        current.source_end_row = excel_row

    if current is not None:
        _validate_order_totals(current)
        orders.append(current)
    if not orders:
        raise ReportParseError("Không tìm thấy đơn hàng hợp lệ trong báo cáo.")
    return orders


def _new_order_from_row(
    workbook: xlrd.book.Book,
    sheet: Sheet,
    *,
    row_index: int,
    context: Sequence[str],
    report_start: date,
    report_end: date,
) -> OrderRecord:
    excel_row = row_index + 1
    required_columns = (0, 1, 2, 3, 4, 5, 7, 8, 9, 10)
    for column in required_columns:
        if not context[column]:
            raise ReportParseError(
                f"Thiếu giá trị bắt buộc '{EXPECTED_HEADERS[column]}'.",
                cell=f"{_column_letter(column)}{excel_row}",
            )

    sale_date = _date_from_cell(
        workbook, sheet.cell(row_index, 1), field_name="Ngày", excel_row=excel_row
    )
    invoice_date = _date_from_cell(
        workbook,
        sheet.cell(row_index, 5),
        field_name="Ngày hóa đơn",
        excel_row=excel_row,
    )
    related_order_date = _date_from_cell(
        workbook,
        sheet.cell(row_index, 8),
        field_name="Ngày đặt đơn hàng liên quan",
        excel_row=excel_row,
    )

    for field_name, value in (("Ngày", sale_date), ("Ngày hóa đơn", invoice_date)):
        if not report_start <= value <= report_end:
            raise ReportParseError(
                f"{field_name} {_format_date(value)} nằm ngoài khoảng báo cáo "
                f"{_format_date(report_start)} - {_format_date(report_end)}.",
                cell=f"{'B' if field_name == 'Ngày' else 'F'}{excel_row}",
            )
    if related_order_date > report_end:
        raise ReportParseError(
            "Ngày đặt đơn hàng liên quan nằm sau ngày báo cáo.",
            cell=f"I{excel_row}",
        )

    sequence_decimal = _required_decimal(
        sheet.cell(row_index, 0), field_name="STT", excel_row=excel_row
    )
    if sequence_decimal != sequence_decimal.to_integral_value():
        raise ReportParseError("STT phải là số nguyên.", cell=f"A{excel_row}")

    bill_pct, bill_pct_error = _optional_decimal_or_error(
        sheet.cell(row_index, 21), field_name="Chiết khấu tổng bill (%)"
    )
    bill_amount, bill_amount_error = _optional_decimal_or_error(
        sheet.cell(row_index, 22), field_name="Tiền chiết khấu tổng bill"
    )

    return OrderRecord(
        sequence=int(sequence_decimal),
        source_start_row=excel_row,
        source_end_row=excel_row,
        sale_date=sale_date,
        employee_code=context[2],
        employee_name=_clean_display_text(context[3]),
        invoice_number=context[4],
        invoice_date=invoice_date,
        salesperson=_clean_display_text(context[6]),
        related_order_code=context[7],
        related_order_date=related_order_date,
        customer_code=context[9],
        customer_name=_clean_display_text(context[10]),
        customer_address=_clean_display_text(context[11]),
        bill_discount_percent=bill_pct,
        bill_discount_amount=bill_amount,
        bill_discount_percent_error=bill_pct_error,
        bill_discount_amount_error=bill_amount_error,
        payment_amount=_required_decimal(
            sheet.cell(row_index, 23), field_name="Thanh toán", excel_row=excel_row
        ),
        revenue_ex_vat=_required_decimal(
            sheet.cell(row_index, 24),
            field_name="Doanh thu (-VAT)",
            excel_row=excel_row,
        ),
        note=_clean_display_text(_cell_text(sheet.cell(row_index, 25))),
    )


def _forward_fill_and_validate_context(
    raw_context: Sequence[str],
    previous_context: Sequence[str],
    *,
    row_number: int,
) -> list[str]:
    result = list(previous_context)
    for column, raw_value in enumerate(raw_context):
        if not raw_value:
            continue
        if column == 7:
            # A non-empty H is handled as a new group before this function.
            continue
        if raw_value != previous_context[column]:
            raise ReportParseError(
                f"Giá trị '{EXPECTED_HEADERS[column]}' thay đổi giữa cùng một đơn.",
                cell=f"{_column_letter(column)}{row_number}",
                hint="Kiểm tra lại cấu trúc nhóm trong file Sopoka.",
            )
        result[column] = raw_value
    return result


def _parse_item(workbook: xlrd.book.Book, sheet: Sheet, row_index: int) -> OrderItem:
    excel_row = row_index + 1
    product_code = _identifier(
        workbook, sheet.cell(row_index, 12), field_name="Mã sản phẩm"
    )
    product_name = _clean_display_text(_cell_text(sheet.cell(row_index, 13)))
    category = _clean_display_text(_cell_text(sheet.cell(row_index, 14)))
    unit = _clean_display_text(_cell_text(sheet.cell(row_index, 17)))
    for column, value in ((12, product_code), (13, product_name), (14, category), (17, unit)):
        if not value:
            raise ReportParseError(
                f"Thiếu giá trị bắt buộc '{EXPECTED_HEADERS[column]}'.",
                cell=f"{_column_letter(column)}{excel_row}",
            )

    quantity = _required_decimal(
        sheet.cell(row_index, 15), field_name="Số lượng bán", excel_row=excel_row
    )
    promotion_quantity = _required_decimal(
        sheet.cell(row_index, 16), field_name="Số lượng KM", excel_row=excel_row
    )
    gross_amount = _required_decimal(
        sheet.cell(row_index, 18),
        field_name="Thành tiền (SL bán x giá)",
        excel_row=excel_row,
    )
    discount_percent = _required_decimal(
        sheet.cell(row_index, 19),
        field_name="Chiết khấu (%) (từng SP)",
        excel_row=excel_row,
    )
    discount_amount = _required_decimal(
        sheet.cell(row_index, 20),
        field_name="Tiền chiết khấu (từng SP)",
        excel_row=excel_row,
    )

    if quantity <= 0:
        raise ReportParseError("Số lượng bán phải lớn hơn 0.", cell=f"P{excel_row}")
    if promotion_quantity < 0:
        raise ReportParseError("Số lượng khuyến mãi không được âm.", cell=f"Q{excel_row}")
    if gross_amount < 0:
        raise ReportParseError("Thành tiền không được âm.", cell=f"S{excel_row}")
    if not Decimal("0") <= discount_percent <= Decimal("100"):
        raise ReportParseError(
            "Phần trăm chiết khấu phải nằm trong khoảng 0..100.",
            cell=f"T{excel_row}",
        )
    if discount_amount < 0 or discount_amount > gross_amount:
        raise ReportParseError(
            "Tiền chiết khấu dòng không hợp lệ.",
            cell=f"U{excel_row}",
        )

    expected_discount = gross_amount * discount_percent / Decimal("100")
    if not _close_money(expected_discount, discount_amount):
        raise ReportParseError(
            f"Tiền chiết khấu dòng không khớp: tính được "
            f"{_decimal_label(expected_discount)}, file là {_decimal_label(discount_amount)}.",
            cell=f"U{excel_row}",
        )

    unit_price = gross_amount / quantity
    return OrderItem(
        source_row=excel_row,
        product_code=product_code,
        product_name=product_name,
        category=category,
        quantity=quantity,
        promotion_quantity=promotion_quantity,
        unit=unit,
        gross_amount=gross_amount,
        unit_price=unit_price,
        discount_percent=discount_percent,
        discount_amount=discount_amount,
        net_amount=gross_amount - discount_amount,
    )


def _validate_order_totals(order: OrderRecord) -> None:
    if not order.items:
        raise ReportParseError(
            f"Đơn {order.related_order_code} không có sản phẩm.",
            cell=f"H{order.source_start_row}",
        )
    if not _close_money(order.calculated_net_total, order.payment_amount):
        raise ReportParseError(
            f"Tổng đơn {order.related_order_code} không khớp: tổng dòng là "
            f"{_decimal_label(order.calculated_net_total)}, cột Thanh toán là "
            f"{_decimal_label(order.payment_amount)}.",
            cell=f"X{order.source_start_row}",
        )
    if not _close_money(order.payment_amount, order.revenue_ex_vat):
        raise ReportParseError(
            f"Thanh toán và Doanh thu (-VAT) của đơn {order.related_order_code} "
            "không khớp.",
            cell=f"Y{order.source_start_row}",
        )


def _ensure_unique_business_keys(orders: Sequence[OrderRecord]) -> None:
    seen_orders: dict[str, int] = {}
    seen_invoices: dict[str, int] = {}
    for order in orders:
        if order.related_order_code in seen_orders:
            raise ReportParseError(
                f"Mã đặt đơn hàng '{order.related_order_code}' bị lặp trong file.",
                cell=f"H{order.source_start_row}",
                hint=f"Mã này đã xuất hiện ở dòng {seen_orders[order.related_order_code]}.",
            )
        if order.invoice_number in seen_invoices:
            raise ReportParseError(
                f"Số hóa đơn '{order.invoice_number}' bị lặp trong file.",
                cell=f"E{order.source_start_row}",
                hint=f"Hóa đơn đã xuất hiện ở dòng {seen_invoices[order.invoice_number]}.",
            )
        seen_orders[order.related_order_code] = order.source_start_row
        seen_invoices[order.invoice_number] = order.source_start_row


def _build_qa(orders: Sequence[OrderRecord], *, batch_count: int) -> ReportQA:
    gross_total = sum((order.gross_total for order in orders), Decimal("0"))
    discount_total = sum(
        (order.line_discount_total for order in orders), Decimal("0")
    )
    net_total = sum(
        (order.calculated_net_total for order in orders), Decimal("0")
    )
    payment_total = sum((order.payment_amount for order in orders), Decimal("0"))
    revenue_total = sum((order.revenue_ex_vat for order in orders), Decimal("0"))
    checks = {
        "line_discount_math_ok": all(
            _close_money(
                item.gross_amount * item.discount_percent / Decimal("100"),
                item.discount_amount,
            )
            for order in orders
            for item in order.items
        ),
        "order_payment_math_ok": all(
            _close_money(order.calculated_net_total, order.payment_amount)
            for order in orders
        ),
        "payment_equals_revenue_ok": all(
            _close_money(order.payment_amount, order.revenue_ex_vat)
            for order in orders
        ),
        "unique_related_order_codes_ok": len(
            {order.related_order_code for order in orders}
        )
        == len(orders),
        "unique_invoice_numbers_ok": len({order.invoice_number for order in orders})
        == len(orders),
    }
    if not all(checks.values()):
        failed = ", ".join(name for name, passed in checks.items() if not passed)
        raise ReportParseError(f"Kiểm tra QA tổng hợp thất bại: {failed}.")
    return ReportQA(
        order_count=len(orders),
        line_count=sum(len(order.items) for order in orders),
        gross_total=gross_total,
        line_discount_total=discount_total,
        calculated_net_total=net_total,
        payment_total=payment_total,
        revenue_ex_vat_total=revenue_total,
        batch_count=batch_count,
        checks=checks,
    )


def _date_from_cell(
    workbook: xlrd.book.Book,
    cell: Cell,
    *,
    field_name: str,
    excel_row: int,
) -> date:
    cell_ref = f"{_column_letter(_header_index(field_name))}{excel_row}"
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, workbook.datemode).date()
        except (ValueError, OverflowError) as exc:
            raise ReportParseError(f"{field_name} không hợp lệ.", cell=cell_ref) from exc
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        try:
            return xlrd.xldate_as_datetime(cell.value, workbook.datemode).date()
        except (ValueError, OverflowError) as exc:
            raise ReportParseError(f"{field_name} không hợp lệ.", cell=cell_ref) from exc
    text = _cell_text(cell)
    if not text:
        raise ReportParseError(f"Thiếu {field_name}.", cell=cell_ref)
    return _parse_date_text(text, field_name=field_name, cell=cell_ref)


def _parse_date_text(
    value: str,
    *,
    field_name: str,
    cell: str | None = None,
) -> date:
    normalized = value.strip()
    for date_format in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(normalized, date_format).date()
        except ValueError:
            continue
    raise ReportParseError(
        f"{field_name} '{value}' không đúng định dạng ngày.",
        cell=cell,
        hint="Định dạng hợp lệ là dd/mm/YYYY.",
    )


def _coerce_expected_date(value: DateInput) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        return _parse_date_text(value, field_name="Ngày yêu cầu")
    raise ReportParseError("Ngày yêu cầu không hợp lệ.")


def _required_decimal(
    cell: Cell,
    *,
    field_name: str,
    excel_row: int,
    column_index: int | None = None,
) -> Decimal:
    column = _header_index(field_name) if column_index is None else column_index
    cell_ref = f"{_column_letter(column)}{excel_row}"
    if cell.ctype == xlrd.XL_CELL_ERROR:
        raise ReportParseError(
            f"{field_name} đang chứa lỗi Excel {_excel_error(cell)}.", cell=cell_ref
        )
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        raise ReportParseError(f"Thiếu {field_name}.", cell=cell_ref)
    try:
        return _to_decimal(cell.value)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ReportParseError(
            f"{field_name} phải là số, nhận '{_cell_text(cell)}'.", cell=cell_ref
        ) from exc


def _optional_decimal_or_error(
    cell: Cell, *, field_name: str
) -> tuple[Decimal | None, str | None]:
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return None, _excel_error(cell)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None, None
    # Sopoka exports blank bill-level discount cells as numeric NaN values in
    # some .xls files.  They are optional fields, so normalize that export
    # quirk to an actual blank rather than rejecting an otherwise valid order.
    if isinstance(cell.value, float) and math.isnan(cell.value):
        return None, None
    try:
        return _to_decimal(cell.value), None
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ReportParseError(f"{field_name} không phải là số hợp lệ.") from exc


def _identifier(
    workbook: xlrd.book.Book,
    cell: Cell,
    *,
    field_name: str,
) -> str:
    """Read an identifier without coercing text to a number.

    Text values keep leading zeroes.  For genuinely numeric identifier cells,
    a zero-only Excel number format (for example ``000000``) is honored.
    """

    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if cell.ctype == xlrd.XL_CELL_ERROR:
        raise ReportParseError(
            f"{field_name} đang chứa lỗi Excel {_excel_error(cell)}."
        )
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        numeric = _to_decimal(cell.value)
        if numeric == numeric.to_integral_value():
            raw = str(int(numeric))
            width = _numeric_zero_pad_width(workbook, cell)
            return raw.zfill(width) if width else raw
        return format(numeric.normalize(), "f")
    return _cell_text(cell).strip()


def _numeric_zero_pad_width(workbook: xlrd.book.Book, cell: Cell) -> int:
    openxml_format = getattr(cell, "number_format", "")
    if openxml_format:
        first_section = openxml_format.split(";", 1)[0]
        first_section = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", first_section)
        groups = re.findall(r"0+", first_section)
        if groups:
            return max(len(group) for group in groups)
    try:
        xf = workbook.xf_list[cell.xf_index]
        format_string = workbook.format_map[xf.format_key].format_str
    except (AttributeError, IndexError, KeyError, TypeError):
        return 0
    first_section = format_string.split(";", 1)[0]
    first_section = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", first_section)
    groups = re.findall(r"0+", first_section)
    return max((len(group) for group in groups), default=0)


def _cell_text(cell: Cell) -> str:
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return _excel_error(cell)
    value = cell.value
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _excel_error(cell: Cell) -> str:
    try:
        code = int(cell.value)
    except (TypeError, ValueError):
        return "#ERROR!"
    return xlrd.error_text_from_code.get(code, f"#ERROR({code})!")


def _to_decimal(value: Any) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool) or value is None:
        raise InvalidOperation
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise InvalidOperation
        return Decimal(str(value))
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if not text:
        raise InvalidOperation
    # Vietnamese display: 1.234.567,89.  Also accept invariant/US numbers.
    if re.fullmatch(r"[-+]?\d{1,3}(?:\.\d{3})+(?:,\d+)?", text):
        text = text.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?", text):
        text = text.replace(",", "")
    elif "," in text and "." not in text:
        text = text.replace(",", ".")
    return Decimal(text)


def _row_has_any_value(sheet: Sheet, row_index: int, column_count: int) -> bool:
    return any(
        sheet.cell(row_index, column).ctype
        not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
        and _cell_text(sheet.cell(row_index, column)) != ""
        for column in range(min(column_count, sheet.ncols))
    )


def _normalize_header(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip()).casefold()


def _clean_display_text(value: str) -> str:
    # Remove accidental leading/trailing whitespace while retaining intentional
    # internal spacing and all Vietnamese Unicode characters.
    return value.strip()


def _header_index(field_name: str) -> int:
    aliases = {
        "Chiết khấu tổng bill (%)": 21,
        "Tiền chiết khấu tổng bill": 22,
    }
    if field_name in aliases:
        return aliases[field_name]
    try:
        return EXPECTED_HEADERS.index(field_name)
    except ValueError as exc:
        raise RuntimeError(f"Unknown internal field name: {field_name}") from exc


def _column_letter(zero_based_column: int) -> str:
    value = zero_based_column + 1
    letters = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _close_money(left: Decimal, right: Decimal) -> bool:
    return abs(left - right) <= MONEY_TOLERANCE


def _format_date(value: date) -> str:
    return value.strftime("%d/%m/%Y")


def _decimal_label(value: Decimal) -> str:
    integral = value.quantize(Decimal("1"))
    return f"{integral:,}".replace(",", ".")


def _json_decimal(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def _make_json_serializable(value: Any) -> Any:
    if isinstance(value, Decimal):
        return _json_decimal(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: _make_json_serializable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_make_json_serializable(item) for item in value]
    return value


__all__ = [
    "EXPECTED_HEADERS",
    "MAX_ORDERS_PER_BATCH",
    "SUPPORTED_REPORT_EXTENSIONS",
    "SUPPORTED_REPORT_FORMATS",
    "OrderBatch",
    "OrderItem",
    "OrderRecord",
    "ParsedReport",
    "ReportParseError",
    "ReportQA",
    "batch_orders",
    "parse_report",
    "parse_report_to_dict",
    "parse_today_report",
]
